"""Build the stage 3 SmartRail poster in the course template (67 x 97 cm).

Run from anywhere:  python Poster/build_poster.py
Writes Poster/SmartRail_Poster.pptx; export the PDF from PowerPoint.
Chart values come from simulation/boarding_sim.py
run with its default parameters.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "input" / "course-poster-template.pptx"
OUT = HERE / "SmartRail_Poster.pptx"
FIGURES = HERE.parent / "simulation" / "figures"
IMAGES = {
    "platform_map.png": FIGURES / "platform_map.png",
    "phone_dispatcher.png": FIGURES / "phone_dispatcher.png",
    "photo_sensor_seats.jpg": HERE / "input" / "stage3" / "photo_sensor_seats.jpg",
}

TEAL = RGBColor(0x08, 0x7A, 0x78)
TEAL_DARK = RGBColor(0x05, 0x5E, 0x5C)
TEAL_TINT = RGBColor(0xE6, 0xF3, 0xF2)
ORANGE = RGBColor(0xD9, 0x6C, 0x1E)
INK = RGBColor(0x23, 0x1F, 0x20)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
LINE = RGBColor(0xC9, 0xDB, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation(TEMPLATE)
slide = prs.slides[0]
shapes = {s.name: s for s in slide.shapes}

# Remove template instructions / placeholders we replace with our own shapes.
for name in ("Rectangle 1", "object 3", "TextBox 19"):
    el = shapes[name]._element
    el.getparent().remove(el)


def run(p, text, size, bold=False, color=INK, italic=False):
    r = p.add_run()
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    return r


def clear(tf):
    """Leave one empty paragraph so runs can be added from scratch."""
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    return p


def textbox(x, y, w, h, margin=0):
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(margin)
    tf.margin_top = tf.margin_bottom = Pt(0)
    return tb, tf


def bullet(p, indent_pt=48):
    """Bullet character on a paragraph, same style as the template list."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent_pt))))
    pPr.set("indent", str(-int(Pt(indent_pt))))
    from lxml import etree
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    clr = etree.SubElement(pPr, f"{{{ns}}}buClr")
    etree.SubElement(clr, f"{{{ns}}}srgbClr", val="087A78")
    etree.SubElement(pPr, f"{{{ns}}}buFont", typeface="Arial")
    etree.SubElement(pPr, f"{{{ns}}}buChar", char="•")


# ---------- Logo ----------
logo = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(190), Pt(372), Pt(215), Pt(215))
logo.name = "Project logo"
logo.adjustments[0] = 0.18
logo.fill.solid()
logo.fill.fore_color.rgb = TEAL
logo.line.fill.background()
logo.shadow.inherit = False
tf = logo.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
run(p, "SR", 120, bold=True, color=WHITE)

# ---------- Name + tagline (template boxes) ----------
p = clear(shapes["TextBox 24"].text_frame)
p.alignment = PP_ALIGN.CENTER
run(p, "SmartRail", 120, bold=True, color=INK)
name_box = shapes["TextBox 24"]
name_box.top = Pt(360)

p = clear(shapes["TextBox 18"].text_frame)
p.alignment = PP_ALIGN.CENTER
run(p, "Spread out. Board faster.", 60, bold=True, color=TEAL)
shapes["TextBox 18"].top = Pt(512)

# ---------- Description (template box, 54pt) ----------
desc = shapes["TextBox 25"]
desc.left, desc.top, desc.width = Pt(140), Pt(632), Pt(1620)
tf = desc.text_frame
p = clear(tf)
D = 54


def spaced(p, before=0):
    p.space_before = Pt(before)
    p.space_after = Pt(0)
    p.line_spacing = 0.95

spaced(p)
run(p, "The problem: ", D, bold=True, color=ORANGE)
run(p, "passengers bunch near the entrances, so a few doors take most boarders and hold the train.", D)

p = tf.add_paragraph(); spaced(p, 14)
run(p, "SmartRail ", D, bold=True, color=TEAL)
run(p, "senses free train seats and shows them on the platform, so passengers spread out before the train arrives.", D)

for label, text in (
    ("Seat sensing: ", "ESP32 + IR sensor above each seat"),
    ("Platform map: ", "each car’s free seats at its zone"),
    ("Phone app: ", "a different car for each passenger"),
):
    p = tf.add_paragraph(); spaced(p, 6)
    bullet(p)
    run(p, label, D, bold=True, color=TEAL_DARK)
    run(p, text, D)

p = tf.add_paragraph(); spaced(p, 14)
run(p, "Result: ", D, bold=True, color=ORANGE)
run(p, "in simulation, stops are 22% shorter.", D)

X0, X1 = 145, 1755


def caption(x, y, w, text, h=50, size=38):
    tb, tf = textbox(x, y, w, h)
    tb.name = "Caption"
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 0.92
    head, _, rest = text.partition(":")
    if rest:
        run(p, head + ":", size, bold=True, color=TEAL_DARK)
        run(p, rest, size, color=INK)
    else:
        run(p, text, size, bold=True, color=TEAL_DARK)


# ---------- System diagram (native, editable shapes) ----------
ARROW_W = 96
BLOCK_W = (X1 - X0 - 3 * ARROW_W) / 4
BY, BH = 1262, 282

blocks = [
    ("Seat sensors", [(None, "IR sensor above each seat"),
                      (None, "2 real seats: A1, B1")]),
    ("ESP32", [("1", "Calibrate"),
               ("2", "Filter and decide"),
               ("3", "Send seat states")]),
    ("Laptop server", [("4", "Add 1,294 simulated seats"),
                       ("5", "Count free seats per car")]),
    ("Passengers", [("6", "Platform map + phone app")]),
]
arrow_labels = ["analog", "USB serial", "Wi-Fi / LAN"]

for i, (title, lines) in enumerate(blocks):
    x = X0 + i * (BLOCK_W + ARROW_W)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(BY), Pt(BLOCK_W), Pt(BH))
    box.name = f"Diagram block {i + 1}"
    box.adjustments[0] = 0.07
    box.fill.solid(); box.fill.fore_color.rgb = TEAL_TINT
    box.line.color.rgb = TEAL; box.line.width = Pt(4)
    box.shadow.inherit = False
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Pt(16)
    tf.margin_top = Pt(12); tf.margin_bottom = Pt(6)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, title, 42, bold=True, color=TEAL_DARK)
    for num, text in lines:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(8)
        p.line_spacing = 0.9
        if num:
            run(p, num + "  ", 40, bold=True, color=ORANGE)
        run(p, text, 36, color=INK)

    if i < 3:
        ax = x + BLOCK_W + 12
        arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Pt(ax), Pt(BY + BH / 2 - 56), Pt(ARROW_W - 24), Pt(60))
        arr.name = f"Diagram arrow {i + 1}"
        arr.fill.solid(); arr.fill.fore_color.rgb = TEAL
        arr.line.fill.background()
        arr.shadow.inherit = False
        tb, tf = textbox(x + BLOCK_W - 4, BY + BH / 2 + 10, ARROW_W + 8, 90)
        tb.name = f"Diagram arrow label {i + 1}"
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.line_spacing = 0.9
        run(p, arrow_labels[i], 28, bold=True, color=MUTED)

caption(X0, BY + BH + 6, X1 - X0, "System diagram: from seat sensor to the passenger, steps 1–6")

# ---------- Simulation results (native charts, data from simulation/boarding_sim.py) ----------
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE

GREY = RGBColor(0xB4, 0xC0, 0xC8)
CW = (X1 - X0 - 40) / 2
CY, CH = 1606, 316


def style_chart(chart, legend=True):
    chart.font.size = Pt(26)
    chart.font.color.rgb = INK
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(28)
    va = chart.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE2, 0xE8, 0xEC)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(24)
    chart.category_axis.tick_labels.font.size = Pt(24)
    chart.category_axis.format.line.color.rgb = GREY


# Mean passengers boarding each carriage over 500 stops (boarders_per_car.csv).
bars = CategoryChartData()
bars.categories = [str(n) for n in range(1, 13)]
bars.add_series("No guidance", (16.67, 30.47, 37.6, 31.8, 21.04, 17.26, 21.92, 24.89, 20.54, 11.28, 4.43, 1.16))
bars.add_series("Platform map", (22.41, 25.86, 27.3, 26.73, 24.58, 22.9, 21.57, 20.78, 17.83, 13.9, 9.63, 5.58))
gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(X0), Pt(CY), Pt(CW), Pt(CH), bars)
gf.name = "Chart boarders per car"
chart = gf.chart
style_chart(chart)
chart.plots[0].gap_width = 40
chart.plots[0].overlap = 0
for series, color in zip(chart.series, (GREY, TEAL)):
    series.format.fill.solid(); series.format.fill.fore_color.rgb = color
chart.value_axis.maximum_scale = 40
chart.value_axis.major_unit = 10
caption(X0, CY + CH + 6, CW, "Boarders per car: the crowd spreads out")

# Mean dwell against compliance (compliance_sweep.csv; ranked display from sweep_compliance(top3)).
lines = CategoryChartData()
lines.categories = [f"{n}%" if n % 20 == 0 else "" for n in range(0, 101, 10)]
lines.add_series("Each car shows its own state",
                 (142.0, 130.4, 123.1, 117.5, 114.7, 112.2, 110.8, 110.2, 109.7, 109.2, 108.0))
lines.add_series("Names the 3 best cars",
                 (142.0, 138.5, 136.0, 143.9, 158.4, 178.4, 200.9, 229.0, 259.0, 288.6, 322.5))
gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Pt(X0 + CW + 40), Pt(CY), Pt(CW), Pt(CH), lines)
gf.name = "Chart dwell vs compliance"
chart = gf.chart
style_chart(chart)
chart.value_axis.minimum_scale = 100
chart.value_axis.maximum_scale = 330
chart.value_axis.major_unit = 50
for series, color in zip(chart.series, (TEAL, ORANGE)):
    series.smooth = False
    series.format.line.color.rgb = color
    series.format.line.width = Pt(5)
    series.marker.style = XL_MARKER_STYLE.CIRCLE
    series.marker.size = 10
    series.marker.format.fill.solid(); series.marker.format.fill.fore_color.rgb = color
    series.marker.format.line.color.rgb = color
caption(X0 + CW + 40, CY + CH + 6, CW, "Stop time (s) vs. share who follow it")

# ---------- Interfaces and prototype ----------
IY, IH = 1990, 262
photos = [
    ("platform_map.png", 1920 / 1080, "Platform screen: no ranking"),
    ("phone_dispatcher.png", 1700 / 1858, "Phone app: one car each"),
    ("photo_sensor_seats.jpg", 2048 / 1077, "Seat sensor prototype"),
]
widths = [IH * r for _, r, _ in photos]
gap = (X1 - X0 - sum(widths)) / (len(photos) - 1)
x = X0
for (fname, ratio, cap), w in zip(photos, widths):
    pic = slide.shapes.add_picture(str(IMAGES[fname]), Pt(x), Pt(IY), Pt(w), Pt(IH))
    pic.name = cap
    pic.line.color.rgb = LINE
    pic.line.width = Pt(2)
    cw = max(w, 470)
    caption(x + w / 2 - cw / 2, IY + IH + 8, cw, cap)
    x += w + gap

# ---------- Team (template box) ----------
team = shapes["TextBox 27"]
team.top = Pt(2392)
tf = team.text_frame
p = clear(tf)
p.alignment = PP_ALIGN.CENTER
run(p, "Team: ", 50, bold=True, color=TEAL_DARK)
run(p, "Yazan Rabea  •  Abed Azem  •  Samar Khatib", 50, bold=True, color=INK)
p = tf.add_paragraph()
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(10)
run(p, "Instructors: ", 44, bold=True, color=TEAL_DARK)
run(p, "Itai Dabran  •  Diana Cohen  •  Ido Ram", 44, color=INK)

prs.save(OUT)
print("saved", OUT)
